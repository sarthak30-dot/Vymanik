#!/usr/bin/env python3
"""
Regenerate the _rawAnomalies array in src/lib/mock-data.ts from the client's
corrected March 2026 deliverable: an xlsx with the authoritative Defect Type /
Severity (COA1-3) / Delta_T per panel, joined against the original survey KML
for the footprint geometry and GPS centroid that xlsx doesn't carry.

WHY A JOIN AND NOT JUST THE XLSX OR JUST THE KML
-------------------------------------------------
The xlsx is the corrected source of truth for `type`/`severity`/`deltaT` — see
below. But it has no polygon geometry, only a centroid (xcoord/ycoord), and the
map needs the real per-panel outline (see PANEL_FILL_MIN_ZOOM in
routes/_app/map.tsx). The KML has that geometry but — this is the reason this
script exists rather than reusing import_defect_kml.py — its defect code's
meaning is NOT the one that script's CODE_MAP documents. Same two-letter codes
(MM, MC, SP, BD, SO, MB, SC, SH), different real-world meaning this time (e.g.
MM here is "Module Missing", not "Multi-Module Hotspot"). Decoding the KML's
`id` field with the old table would relabel 578 "Module Missing" panels as
"Multi-Module Hotspot" and be flatly wrong, which is exactly the bug this
script fixes — so `type`/`severity`/`deltaT` always come from the xlsx, never
from the KML's code.

JOIN KEY
--------
(block, inv, smb, table, str, rack, panel) — verified 1:1 against both files
before trusting it: 1249/1249 xlsx rows matched, 1249/1249 KML placemarks
matched, zero orphans either side. Do not join on Label (image filename):
multiple panels share one capture frame, so it's not unique (484 distinct
labels across the 1249 rows).

USAGE
    python3 scripts/import_corrected_defect_survey.py \
        "DFECT PANELS corrected.xlsx" defpanels1.kml > /tmp/anomalies.ts

Then splice the output into mock-data.ts's `const _rawAnomalies = ([ ... ]`
block by hand (or adapt gen_mock_from_csv.py's regex-replace approach) —
and recompute plant.healthScore / dailyLossINR / dailyLossKWh / allPlants[0]'s
criticalCount+mediumCount / inspectionHistory's row from the new severity mix,
since none of those are derived automatically from _rawAnomalies.
"""
import sys
import xml.etree.ElementTree as ET
from collections import Counter

import openpyxl

INSPECTION_DATE = "31 March 2026"

SEV_MAP = {"COA3": "critical", "COA2": "medium", "COA1": "normal"}
SEV_RANK = {"critical": 0, "medium": 1, "normal": 2}

K = "{http://www.opengis.net/kml/2.2}"


def num(x: float) -> str:
    """Trim float noise without losing sub-centimetre precision (~7 decimals)."""
    return f"{x:.7f}".rstrip("0").rstrip(".")


def norm_block(v) -> str:
    """'006' / '6.0' / '6' all collapse to '6' — see import_defect_kml.py's
    same normalisation note on inconsistent zero-padding."""
    s = str(v).strip()
    return str(int(float(s))) if s.replace(".", "", 1).isdigit() else s


def norm_smb(v) -> str:
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s.zfill(2) if s.isdigit() else s


def key_of(block, inv, smb, table, strside, rack, panel):
    return (norm_block(block), str(inv).strip(), norm_smb(smb),
            str(int(float(table))), str(strside).strip(),
            str(int(float(rack))), str(int(float(panel))))


def load_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Copy of DFECT PANELS"]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    by_key = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None and all(v is None for v in r):
            continue
        k = key_of(r[idx["Block"]], r[idx["Inv"]], r[idx["SMB"]], r[idx["Table"]],
                   r[idx["String"]], r[idx["Rack"]], r[idx["Panel"]])
        by_key[k] = {
            "defect_type": str(r[idx["Defect Type"]]).strip(),
            "severity_coa": str(r[idx["Severity"]]).strip(),
            "delta_t": float(r[idx["Delta_T"]]),
            "label": str(r[idx["Label"]]).strip(),
            "module": str(r[idx["Module"]]).strip(),
        }
    return by_key


def load_kml(path):
    tree = ET.parse(path)
    out = []
    for pm in tree.iter(K + "Placemark"):
        data = {sd.get("name"): (sd.text or "").strip() for sd in pm.iter(K + "SimpleData")}
        ring_el = pm.find(f".//{K}LinearRing/{K}coordinates")
        if ring_el is None or not data:
            continue
        pts = []
        for tok in ring_el.text.split():
            lng, lat = tok.split(",")[:2]
            pts.append((float(lng), float(lat)))
        out.append((data, pts))
    return out


def main(xlsx_path, kml_path):
    xlsx_by_key = load_xlsx(xlsx_path)
    rows, missing = [], 0

    for data, pts in load_kml(kml_path):
        uniq = pts[:-1] if pts[0] == pts[-1] else pts
        clng = sum(p[0] for p in uniq) / len(uniq)
        clat = sum(p[1] for p in uniq) / len(uniq)

        k = key_of(data.get("block"), data.get("inv"), data.get("smb"), data.get("table"),
                   data.get("str"), data.get("Rack"), data.get("Panel"))
        x = xlsx_by_key.get(k)
        if x is None:
            missing += 1
            continue

        rows.append({
            "rack": str(int(float(data.get("Rack")))),
            "panel": str(int(float(data.get("Panel")))),
            "module": x["module"] or data.get("Module", ""),
            "block": norm_block(data.get("block")),
            "inv": str(data.get("inv")).strip(),
            "smb": norm_smb(data.get("smb")),
            "table": str(int(float(data.get("table")))),
            "str": str(data.get("str")).strip(),
            "code": data.get("id", "").upper(),
            "type": x["defect_type"],
            "severity": SEV_MAP[x["severity_coa"]],
            "delta_t": x["delta_t"],
            "label": x["label"] or data.get("Label", ""),
            "lat": clat, "lng": clng, "ring": pts,
        })

    print(f"// matched {len(rows)} / kml-only {missing} / xlsx-only {len(xlsx_by_key) - len(rows)}",
          file=sys.stderr)
    if missing or len(rows) != len(xlsx_by_key):
        print("// WARNING: join is not a clean 1:1 — verify before trusting this output", file=sys.stderr)

    rows.sort(key=lambda r: (SEV_RANK[r["severity"]], r["type"], int(r["rack"]), int(r["panel"])))

    print(f"// {len(rows)} defects — generated by scripts/import_corrected_defect_survey.py")
    for i, r in enumerate(rows, 1):
        ring = ",".join(f"[{num(x)},{num(y)}]" for x, y in r["ring"])
        print(
            f'  {{ id: "{i}", panelId: "R{r["rack"]}-P{r["panel"]}", '
            f'row: {r["rack"]}, col: {r["panel"]}, '
            f'type: "{r["type"]}", deltaT: {r["delta_t"]}, severity: "{r["severity"]}", '
            f'string: "Table-{r["table"]}", inverter: "INV-{r["inv"]}", '
            f'status: "New", date: "{INSPECTION_DATE}", inspectionTime: "—", '
            f'rgbNote: "Image: {r["label"]}", '
            f'gps: {{ lat: {num(r["lat"])}, lng: {num(r["lng"])} }}, '
            f'block: "{r["block"]}", smb: "{r["smb"]}", stringSide: "{r["str"]}", '
            f'module: "{r["module"]}", defectCode: "{r["code"]}", '
            f'footprint: [{ring}] }},'
        )

    print(f"// Severity: {dict(Counter(r['severity'] for r in rows).most_common())}", file=sys.stderr)
    print(f"// Types:    {dict(Counter(r['type'] for r in rows).most_common())}", file=sys.stderr)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
